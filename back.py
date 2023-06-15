import os
from flask import Flask, render_template, request, send_file
from docx import Document
from docx.shared import Pt

app = Flask(__name__)

def generate_document(field1, field2):


    # Load the document template
    doc = Document('dov.docx')

    # Replace placeholders with form values
    for paragraph in doc.paragraphs:
        if '{{ field1 }}' in paragraph.text:
            run = paragraph.runs[0]
            font = run.font
            font.size = Pt(12)  # Adjust font size if needed
            paragraph.text = paragraph.text.replace('{{ field1 }}', field1)

        if '{{ field2 }}' in paragraph.text:
            run = paragraph.runs[0]
            font = run.font
            font.size = Pt(12)  # Adjust font size if needed
            paragraph.text = paragraph.text.replace('{{ field2 }}', field2)

        # Add more replacements as needed

    # Save the generated document
    output_file =  'generated_document.docx'
    doc.save(output_file)

    return output_file

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate_document', methods=['POST'])
def generate_docx():
    field1 = request.form['field1']
    field2 = request.form['field2']
    # Add more form fields as needed

    generated_doc = generate_document(field1, field2)
    return send_file(generated_doc, as_attachment=True)

if __name__ == '__main__':
    app.run()
